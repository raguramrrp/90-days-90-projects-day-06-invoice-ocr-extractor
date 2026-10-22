"""Unit tests for Excel (.xlsx) Exporter."""

from __future__ import annotations

import io
import openpyxl
import pytest

from src.exporter.excel_exporter import export_to_excel, sanitize_formula_injection


def test_sanitize_formula_injection():
    assert sanitize_formula_injection("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_formula_injection("+12345") == "'+12345"
    assert sanitize_formula_injection("@special") == "'@special"
    assert sanitize_formula_injection("Normal Vendor") == "Normal Vendor"
    assert sanitize_formula_injection(1250.00) == 1250.00


def test_export_to_excel():
    sample_invoices = [
        {
            "id": 1,
            "filename": "inv1.png",
            "vendor_name": "Apex Hardware Labs",
            "invoice_number": "INV-001",
            "invoice_date": "2026-10-15",
            "due_date": "2026-11-15",
            "po_number": "PO-99120",
            "currency": "USD",
            "subtotal": 1000.0,
            "tax_amount": 80.0,
            "shipping_amount": 20.0,
            "discount_amount": 0.0,
            "total_amount": 1100.0,
            "payment_terms": "Net 30",
            "validation_status": "VALIDATED",
            "line_items": [
                {"description": "Monitor", "quantity": 2.0, "unit_price": 500.0, "amount": 1000.0}
            ],
        }
    ]

    excel_bytes = export_to_excel(sample_invoices)
    assert len(excel_bytes) > 0

    # Parse back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Invoices Register" in sheet_names
    assert "Itemized Line Items" in sheet_names
    assert "Audit Summary" in sheet_names

    ws_inv = wb["Invoices Register"]
    assert ws_inv.cell(row=2, column=2).value == "Apex Hardware Labs"
    assert ws_inv.cell(row=2, column=3).value == "INV-001"
