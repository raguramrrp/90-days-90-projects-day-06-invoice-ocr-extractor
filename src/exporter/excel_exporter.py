"""Excel (.xlsx) Workbook Exporter for Invoice Registers and Line Items.

Generates multi-sheet formatted financial workbooks using openpyxl,
protects against spreadsheet formula injection, and formats monetary values.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.validator.duplicate_detector import DuplicateDetector
from src.validator.rule_engine import ValidationRuleEngine


def sanitize_formula_injection(val: Any) -> Any:
    """Prepend single quote if string begins with formula triggers (=, +, -, @, tab, cr)."""
    if isinstance(val, str) and val:
        if val[0] in ("=", "+", "-", "@", "\t", "\r"):
            return f"'{val}"
    return val


def export_to_excel(invoices: List[Dict[str, Any]]) -> bytes:
    """Generate multi-sheet Excel workbook containing invoices, itemized rows, and audit summary."""
    wb = openpyxl.Workbook()

    # Styling definitions
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # =========================================================================
    # SHEET 1: INVOICES REGISTER
    # =========================================================================
    ws_invoices = wb.active
    ws_invoices.title = "Invoices Register"

    headers_invoices = [
        "ID", "Vendor Name", "Invoice #", "PO Number", "Invoice Date", "Due Date",
        "Currency", "Subtotal", "Tax Amount", "Shipping", "Discount", "Total Amount",
        "Payment Terms", "Audit Status", "Processed File", "Recorded At"
    ]
    ws_invoices.append(headers_invoices)

    for col_num in range(1, len(headers_invoices) + 1):
        cell = ws_invoices.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, inv in enumerate(invoices, start=2):
        row_data = [
            inv.get("id"),
            sanitize_formula_injection(inv.get("vendor_name")),
            sanitize_formula_injection(inv.get("invoice_number")),
            sanitize_formula_injection(inv.get("po_number")),
            inv.get("invoice_date"),
            inv.get("due_date"),
            inv.get("currency", "USD"),
            float(inv.get("subtotal", 0.0) or 0.0),
            float(inv.get("tax_amount", 0.0) or 0.0),
            float(inv.get("shipping_amount", 0.0) or 0.0),
            float(inv.get("discount_amount", 0.0) or 0.0),
            float(inv.get("total_amount", 0.0) or 0.0),
            sanitize_formula_injection(inv.get("payment_terms")),
            inv.get("validation_status"),
            sanitize_formula_injection(inv.get("filename")),
            str(inv.get("created_at", "")),
        ]
        ws_invoices.append(row_data)

        # Style row
        use_fill = zebra_fill if row_idx % 2 == 0 else None
        for col_num in range(1, len(headers_invoices) + 1):
            c = ws_invoices.cell(row=row_idx, column=col_num)
            c.border = thin_border
            if use_fill:
                c.fill = use_fill
            # Format numbers
            if col_num in (8, 9, 10, 11, 12):
                c.number_format = "$#,##0.00"
                c.alignment = Alignment(horizontal="right")

    # =========================================================================
    # SHEET 2: ITEMIZED LINE ITEMS
    # =========================================================================
    ws_items = wb.create_sheet(title="Itemized Line Items")
    headers_items = [
        "Invoice ID", "Invoice #", "Vendor Name", "Currency", "Item Description",
        "Quantity", "Unit Price", "Line Total"
    ]
    ws_items.append(headers_items)

    for col_num in range(1, len(headers_items) + 1):
        cell = ws_items.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    item_row_idx = 2
    for inv in invoices:
        inv_id = inv.get("id")
        inv_no = sanitize_formula_injection(inv.get("invoice_number"))
        vendor = sanitize_formula_injection(inv.get("vendor_name"))
        curr = inv.get("currency", "USD")
        items = inv.get("line_items", [])

        if not items:
            ws_items.append([inv_id, inv_no, vendor, curr, "No line items recorded", 1.0, float(inv.get("total_amount", 0.0) or 0.0), float(inv.get("total_amount", 0.0) or 0.0)])
            item_row_idx += 1
        else:
            for itm in items:
                ws_items.append([
                    inv_id,
                    inv_no,
                    vendor,
                    curr,
                    sanitize_formula_injection(itm.get("description", "Item")),
                    float(itm.get("quantity", 1.0) or 1.0),
                    float(itm.get("unit_price", 0.0) or 0.0),
                    float(itm.get("amount", 0.0) or 0.0),
                ])
                # Format numbers
                for col_num in range(1, len(headers_items) + 1):
                    c = ws_items.cell(row=item_row_idx, column=col_num)
                    c.border = thin_border
                    if col_num in (7, 8):
                        c.number_format = "$#,##0.00"
                item_row_idx += 1

    # =========================================================================
    # SHEET 3: AUDIT SUMMARY
    # =========================================================================
    ws_summary = wb.create_sheet(title="Audit Summary")
    ws_summary.append(["Audit Executive Summary", ""])
    ws_summary.append(["Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_summary.append(["Total Invoices Exported", len(invoices)])
    total_spend = sum(float(i.get("total_amount", 0.0) or 0.0) for i in invoices)
    ws_summary.append(["Cumulative Spend ($)", total_spend])
    ws_summary.cell(row=4, column=2).number_format = "$#,##0.00"

    # =========================================================================
    # SHEET 4: VALIDATION EXCEPTIONS
    # =========================================================================
    exceptions = ValidationRuleEngine.evaluate_all(invoices)
    ws_ex = wb.create_sheet(title="Validation Exceptions")
    headers_ex = ["Invoice ID", "Invoice Ref", "Rule ID", "Field", "Severity", "Current Value", "Explanation", "Suggested Action"]
    ws_ex.append(headers_ex)
    for col_num in range(1, len(headers_ex) + 1):
        cell = ws_ex.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for ex_idx, ex in enumerate(exceptions, start=2):
        ws_ex.append([
            ex.invoice_id,
            sanitize_formula_injection(ex.invoice_ref),
            ex.rule_id,
            ex.field,
            ex.severity,
            sanitize_formula_injection(ex.current_value),
            sanitize_formula_injection(ex.explanation),
            sanitize_formula_injection(ex.suggested_action),
        ])
        for col_num in range(1, len(headers_ex) + 1):
            c = ws_ex.cell(row=ex_idx, column=col_num)
            c.border = thin_border

    # =========================================================================
    # SHEET 5: DUPLICATE CANDIDATES
    # =========================================================================
    dup_pairs = DuplicateDetector.find_all_duplicate_pairs(invoices)
    ws_dups = wb.create_sheet(title="Duplicate Review Candidates")
    headers_dup = ["Invoice A ID", "Invoice A #", "Invoice B ID", "Invoice B #", "Vendor", "Match Type", "Decision", "Explanation"]
    ws_dups.append(headers_dup)
    for col_num in range(1, len(headers_dup) + 1):
        cell = ws_dups.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for dup_idx, dp in enumerate(dup_pairs, start=2):
        a = dp["invoice_a"]
        b = dp["invoice_b"]
        ws_dups.append([
            a.get("id"),
            sanitize_formula_injection(a.get("invoice_number")),
            b.get("id"),
            sanitize_formula_injection(b.get("invoice_number")),
            sanitize_formula_injection(a.get("vendor_name")),
            dp.get("match_type"),
            dp.get("decision"),
            sanitize_formula_injection(dp.get("explanation")),
        ])
        for col_num in range(1, len(headers_dup) + 1):
            c = ws_dups.cell(row=dup_idx, column=col_num)
            c.border = thin_border

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_exceptions_to_csv(findings: List[Any]) -> str:
    """Serialize validation findings to CSV format."""
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Invoice ID", "Invoice Ref", "Rule ID", "Field", "Severity", "Current Value", "Explanation", "Suggested Action"])
    for f in findings:
        d = f.to_dict() if hasattr(f, "to_dict") else f
        writer.writerow([
            d.get("invoice_id"),
            sanitize_formula_injection(d.get("invoice_ref")),
            d.get("rule_id"),
            d.get("field"),
            d.get("severity"),
            sanitize_formula_injection(d.get("current_value")),
            sanitize_formula_injection(d.get("explanation")),
            sanitize_formula_injection(d.get("suggested_action")),
        ])
    return output.getvalue()


def export_duplicates_to_csv(pairs: List[Dict[str, Any]]) -> str:
    """Serialize duplicate candidate pairs to CSV format."""
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Invoice A ID", "Invoice A #", "Invoice B ID", "Invoice B #", "Vendor", "Match Type", "Decision", "Explanation"])
    for dp in pairs:
        a = dp.get("invoice_a", {})
        b = dp.get("invoice_b", {})
        writer.writerow([
            a.get("id"),
            sanitize_formula_injection(a.get("invoice_number")),
            b.get("id"),
            sanitize_formula_injection(b.get("invoice_number")),
            sanitize_formula_injection(a.get("vendor_name")),
            dp.get("match_type"),
            dp.get("decision"),
            sanitize_formula_injection(dp.get("explanation")),
        ])
    return output.getvalue()
